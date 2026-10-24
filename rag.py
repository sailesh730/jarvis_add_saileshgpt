"""Extract, index, search, and remove conversation-scoped uploaded documents."""

import os
import uuid
from pathlib import Path

from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_core.runnables import RunnableConfig
from langchain_core.tools import tool
from dotenv import load_dotenv

load_dotenv()

DATA_DIR = Path(__file__).resolve().parent / "data"
VECTOR_DIR = DATA_DIR / "chroma"
COLLECTION_NAME = "uploaded_documents"
CHUNK_SIZE = 1200
CHUNK_OVERLAP = 200


def _get_embeddings():
    provider = os.getenv("RAG_EMBEDDING_PROVIDER", "").strip().lower()
    if not provider:
        provider = "google" if os.getenv("GOOGLE_API_KEY") else "ollama"

    if provider == "google":
        api_key = os.getenv("GOOGLE_API_KEY")
        if not api_key:
            raise RuntimeError(
                "RAG_EMBEDDING_PROVIDER is 'google' but GOOGLE_API_KEY is not configured."
            )
        from langchain_google_genai import GoogleGenerativeAIEmbeddings

        return GoogleGenerativeAIEmbeddings(
            model=os.getenv("RAG_GOOGLE_EMBEDDING_MODEL", "models/gemini-embedding-001"),
            google_api_key=api_key,
        )
    if provider == "ollama":
        from langchain_ollama import OllamaEmbeddings

        model = os.getenv("RAG_OLLAMA_EMBEDDING_MODEL", "nomic-embed-text")
        return OllamaEmbeddings(model=model)
    raise ValueError(
        "Unsupported RAG_EMBEDDING_PROVIDER. Choose 'google' or 'ollama'."
    )


def _vector_store() -> Chroma:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    return Chroma(
        collection_name=COLLECTION_NAME,
        persist_directory=str(VECTOR_DIR),
        embedding_function=_get_embeddings(),
    )


def _extract_text(file_path: Path) -> str:
    extension = file_path.suffix.lower()
    if extension == ".pdf":
        from pypdf import PdfReader

        pages = PdfReader(str(file_path)).pages
        text = "\n\n".join(page.extract_text() or "" for page in pages)
    elif extension == ".docx":
        from docx import Document as WordDocument

        word_document = WordDocument(str(file_path))
        text = "\n".join(paragraph.text for paragraph in word_document.paragraphs)
        for table in word_document.tables:
            text += "\n" + "\n".join(
                " | ".join(cell.text for cell in row.cells) for row in table.rows
            )
    elif extension in {".txt", ".md", ".py", ".csv"}:
        text = file_path.read_text(encoding="utf-8-sig")
    else:
        raise ValueError(f"Unsupported document type: {extension or 'unknown'}.")

    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not text:
        raise ValueError(
            "No extractable text was found. Scanned PDFs need OCR before they can be indexed."
        )
    return text


def _split_text(text: str) -> list[str]:
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        if end < len(text):
            boundary = max(
                text.rfind("\n", start + CHUNK_SIZE // 2, end),
                text.rfind(" ", start + CHUNK_SIZE // 2, end),
            )
            if boundary > start:
                end = boundary
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end == len(text):
            break
        start = max(end - CHUNK_OVERLAP, start + 1)
    return chunks


def add_document_to_rag(
    file_path: str,
    thread_id: str,
    source_name: str | None = None,
) -> dict[str, int]:
    """Extract and persist the uploaded document's chunks for one conversation."""
    path = Path(file_path)
    if not thread_id:
        raise ValueError("Conversation ID is required to index a document.")
    text = _extract_text(path)
    chunks = _split_text(text)
    if not chunks:
        raise ValueError("No text chunks could be created from this document.")

    name = source_name or path.name
    documents = [
        Document(
            page_content=chunk,
            metadata={
                "thread_id": thread_id,
                "source": name,
                "chunk": index,
            },
        )
        for index, chunk in enumerate(chunks)
    ]
    ids = [uuid.uuid4().hex for _ in documents]
    _vector_store().add_documents(documents=documents, ids=ids)
    return {"chunks": len(chunks)}


@tool
def search_uploaded_document(
    query: str,
    config: RunnableConfig,
) -> str:
    """Search only the uploaded documents associated with the current conversation."""
    thread_id = (config.get("configurable") or {}).get("thread_id")
    if not isinstance(thread_id, str) or not thread_id:
        return "Document search failed: conversation context is unavailable."
    query = query.strip()
    if not query:
        return "Document search failed: enter a question to search for."
    results = _vector_store().similarity_search(
        query,
        k=4,
        filter={"thread_id": thread_id},
    )

    if not results:
        return "No matching content was found in documents uploaded to this conversation."
    return "\n\n".join(
        f"Source: {document.metadata.get('source', 'uploaded document')}\n"
        f"{document.page_content}"
        for document in results
    )


def delete_thread_documents(thread_id: str) -> None:
    """Remove all persisted document chunks belonging to one conversation."""
    if not thread_id or not VECTOR_DIR.exists():
        return
    _vector_store().delete(where={"thread_id": thread_id})
